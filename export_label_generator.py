"""EXPORT LABEL 워드 템플릿 파싱 및 라벨 생성"""
import io
import re
from copy import deepcopy

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

LABEL_FONT_NAME = 'Calibri'
SET_CODE_FONT_SIZE_PT = 100
DETAIL_FONT_SIZE_PT = 39

# 콜아크형 11자리 계열: M2122|00001 / M2122I00001 / H3571|00001
# 단축 5자리: Y2991
STOCK_CODE_PATTERN = re.compile(
    r'^(?:[A-Z][\d|Il]{6,}|[A-Z]\d{4})$',
    re.IGNORECASE,
)
PALLET_PATTERN = re.compile(r'pallet\s*:\s*(\d+\s*/\s*\d+)', re.IGNORECASE)
_INVALID_XML_CHARS = re.compile(r'[\x00-\x08\x0B\x0C\x0E-\x1F]')


def _sanitize_docx_text(value):
    text = '' if value is None else str(value)
    return _INVALID_XML_CHARS.sub('', text).strip()


def _normalize_arc_key(code):
    return re.sub(r'[^A-Z0-9]', '', (code or '').upper())


def _element_has_section_break(element):
    if element.tag != qn('w:p'):
        return False
    p_pr = element.find(qn('w:pPr'))
    return p_pr is not None and p_pr.find(qn('w:sectPr')) is not None


def _element_has_page_break(element):
    if _element_has_section_break(element):
        return True
    if element.tag != qn('w:p'):
        return False
    for br in element.findall('.//' + qn('w:br')):
        if br.get(qn('w:type')) == 'page':
            return True
    return False


def _remove_page_breaks_from_element(element):
    if element.tag != qn('w:p'):
        return
    for br in element.findall('.//' + qn('w:br')):
        if br.get(qn('w:type')) == 'page':
            parent = br.getparent()
            if parent is not None:
                parent.remove(br)


def _split_document_pages(doc):
    body = doc.element.body
    children = list(body)
    sect_pr = None
    if children and children[-1].tag == qn('w:sectPr'):
        sect_pr = children.pop()

    pages = [[]]
    for child in children:
        if _element_has_page_break(child):
            clean = deepcopy(child)
            _remove_page_breaks_from_element(clean)
            if _paragraph_has_visible_content(clean):
                pages[-1].append(clean)
            pages.append([])
            continue
        pages[-1].append(deepcopy(child))

    pages = [page for page in pages if page]
    return pages, sect_pr


def _paragraph_has_visible_content(element):
    texts = element.findall('.//' + qn('w:t'))
    return any((t.text or '').strip() for t in texts)


def _doc_from_page_elements(page_elements, sect_pr=None):
    doc = Document()
    body = doc.element.body
    for element in page_elements:
        body.insert(len(body) - 1, deepcopy(element))
    if sect_pr is not None:
        body[-1] = deepcopy(sect_pr)
    return doc


def _extract_pallet_text(doc):
    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        match = PALLET_PATTERN.search(text)
        if match:
            return f"Pallet:{match.group(1).replace(' ', '')}"
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                text = cell.text.strip()
                match = PALLET_PATTERN.search(text)
                if match:
                    return f"Pallet:{match.group(1).replace(' ', '')}"
    return ''


def _parse_stock_rows_from_doc(doc):
    rows = []
    for table in doc.tables:
        stock_col = qty_col = sr_col = desc_col = None
        header_row_idx = None

        for ri, row in enumerate(table.rows):
            cells = [cell.text.strip() for cell in row.cells]
            joined = ' '.join(cells).lower()
            if 'stock code' in joined or ('stock' in joined and 'code' in joined):
                header_row_idx = ri
                for ci, text in enumerate(cells):
                    lower = text.lower()
                    if lower == 'sr':
                        sr_col = ci
                    elif 'stock' in lower and 'code' in lower:
                        stock_col = ci
                    elif 'product' in lower and 'description' in lower:
                        desc_col = ci
                    elif 'quantity' in lower or lower in ('qty', "q'ty"):
                        qty_col = ci
                break

        if header_row_idx is None or stock_col is None:
            continue

        sr_counter = 1
        for row in table.rows[header_row_idx + 1:]:
            cells = [cell.text.strip() for cell in row.cells]
            if len(cells) <= stock_col:
                continue

            stock_code = cells[stock_col].strip()
            if not stock_code:
                continue
            # Word 특수문자 파이프/공백 정규화
            stock_code = (
                stock_code.replace('\u00a0', ' ')
                .replace('｜', '|')
                .replace('Ι', 'I')
                .replace('ｌ', 'l')
                .strip()
            )
            # 셀에 줄바꿈이 있으면 첫 유효 토큰 사용
            if '\n' in stock_code or '\r' in stock_code:
                for token in re.split(r'[\r\n]+', stock_code):
                    token = token.strip()
                    if token and STOCK_CODE_PATTERN.match(token):
                        stock_code = token
                        break
            if stock_code.lower().startswith('pallet') or PALLET_PATTERN.search(stock_code):
                break
            if not STOCK_CODE_PATTERN.match(stock_code):
                continue

            quantity = '1'
            if qty_col is not None and len(cells) > qty_col and cells[qty_col].strip():
                quantity = cells[qty_col].strip()

            sr_value = str(sr_counter)
            if sr_col is not None and len(cells) > sr_col and cells[sr_col].strip():
                sr_value = cells[sr_col].strip()

            product_description = ''
            if desc_col is not None and len(cells) > desc_col:
                product_description = cells[desc_col].strip()

            rows.append({
                'stock_code': stock_code,
                'quantity': quantity,
                'sr': sr_value,
                'product_description': product_description,
            })
            sr_counter += 1

    return rows


def _split_page_into_pallet_parts(page_elements, sect_pr=None):
    """페이지 내 pallet 부분 분리 — 테이블 1개(+ 뒤따르는 Pallet 문단) = pallet 1개"""
    parts = []
    current = []

    def flush():
        nonlocal current
        if not current:
            return
        page_doc = _doc_from_page_elements(current, sect_pr)
        if _parse_stock_rows_from_doc(page_doc):
            parts.append(deepcopy(current))
        current = []

    for element in page_elements or []:
        if element.tag == qn('w:tbl'):
            flush()
            current = [deepcopy(element)]
        elif current:
            current.append(deepcopy(element))
    flush()

    if parts:
        return parts
    if page_elements:
        return [deepcopy(page_elements)]
    return []


def parse_template_entries(filepath):
    """업로드 템플릿을 Stock Code 행 단위 항목 목록으로 파싱 (5자리·11자리 혼재 지원)"""
    doc = Document(filepath)
    pages, sect_pr = _split_document_pages(doc)

    if not pages:
        pages = [[]]

    entries = []
    for page_idx, page_elements in enumerate(pages):
        pallet_parts = _split_page_into_pallet_parts(page_elements, sect_pr)
        for part_idx, part_elements in enumerate(pallet_parts):
            page_doc = _doc_from_page_elements(part_elements, sect_pr)
            table_rows = _parse_stock_rows_from_doc(page_doc)
            if not table_rows:
                continue
            pallet_text = _extract_pallet_text(page_doc)
            multi_page = len(pallet_parts) > 1 or len(pages) > 1
            group_id = f'{page_idx}-{part_idx}'
            # 같은 pallet 테이블의 모든 Stock Code 행 → 각각 라벨 생성
            for idx, row in enumerate(table_rows):
                entries.append({
                    **row,
                    'pallet_text': pallet_text,
                    'page_elements': deepcopy(part_elements),
                    # 원본 테이블(여러 줄) 유지. 그룹 첫 성공 행에서 pallet 페이지 출력
                    'needs_row_update': False,
                    'include_pallet_page': idx == 0,
                    'pallet_group_id': group_id,
                    'from_multi_page': multi_page,
                    'pallet_row_index': idx,
                    'pallet_row_count': len(table_rows),
                })

    return entries


def parse_template_rows(filepath):
    """하위 호환 — Stock Code 행 목록만 반환"""
    return [
        {
            'stock_code': entry['stock_code'],
            'quantity': entry['quantity'],
        }
        for entry in parse_template_entries(filepath)
    ]


def _find_set_code_by_stock_code(cursor, stock_code):
    variants = []
    raw = (stock_code or '').strip()
    if raw:
        variants.extend([
            raw,
            raw.replace('|', 'I'),
            raw.replace('|', 'l'),
            raw.upper(),
            raw.replace('|', 'I').upper(),
        ])

    seen = set()
    for variant in variants:
        if not variant or variant in seen:
            continue
        seen.add(variant)
        cursor.execute(
            '''
            SELECT set_code
            FROM selected_combinations
            WHERE call_arc_code = ?
            LIMIT 1
            ''',
            (variant,),
        )
        row = cursor.fetchone()
        if row and row[0]:
            return row[0]

    target = _normalize_arc_key(stock_code)
    if not target:
        return None

    cursor.execute(
        '''
        SELECT set_code, call_arc_code
        FROM selected_combinations
        WHERE call_arc_code IS NOT NULL AND call_arc_code != ''
        '''
    )
    for set_code, call_arc_code in cursor.fetchall():
        if _normalize_arc_key(call_arc_code) == target:
            return set_code
    return None


def _find_label_line_from_dashboard(cursor, set_code, stock_code):
    cursor.execute(
        '''
        SELECT product_name
        FROM export_quotation_new_products
        WHERE product_code = ?
        ORDER BY updated_at DESC
        LIMIT 1
        ''',
        (set_code,),
    )
    row = cursor.fetchone()
    if not row or not row[0]:
        return None

    product_name = str(row[0]).strip()
    if '_' in product_name:
        prefix = product_name.rsplit('_', 1)[0]
        return f'{prefix}_{stock_code}'
    return f'{product_name}_{stock_code}'


def _build_label_line(product_name, stock_code, product_description=None):
    stock_code = (stock_code or '').strip()
    desc = (product_description or '').strip()
    if desc:
        return f'{desc}_{stock_code}' if stock_code else desc
    product_name = (product_name or '').strip()
    if not product_name:
        return stock_code
    if '_' in product_name:
        prefix = product_name.rsplit('_', 1)[0]
        return f'{prefix}_{stock_code}' if stock_code else prefix
    return f'{product_name}_{stock_code}' if stock_code else product_name


def _find_label_from_part_prices(cursor, stock_code, product_description=None):
    """파트가격 테이블에서 Stock Code 매칭 (Y0149, H3571|00001 등)"""
    raw = (stock_code or '').strip()
    if not raw:
        return None
    keys = []
    for key in (raw, raw.replace('|', 'I'), raw.replace('|', ''), raw.upper()):
        if key and key not in keys:
            keys.append(key)

    for key in keys:
        cursor.execute(
            '''
            SELECT material_code, description
            FROM part_prices
            WHERE UPPER(COALESCE(description, '')) LIKE UPPER(?)
               OR UPPER(COALESCE(description, '')) LIKE UPPER(?)
               OR UPPER(COALESCE(material_code, '')) = UPPER(?)
            ORDER BY
                CASE
                    WHEN UPPER(COALESCE(description, '')) LIKE UPPER(?) THEN 0
                    WHEN UPPER(COALESCE(material_code, '')) = UPPER(?) THEN 1
                    ELSE 2
                END,
                LENGTH(COALESCE(description, '')) DESC,
                updated_at DESC
            LIMIT 1
            ''',
            (f'%_{key}', f'%{key}%', key, f'%_{key}', key),
        )
        row = cursor.fetchone()
        if row and row[0]:
            material_code = str(row[0]).strip()
            description = str(row[1] or '').strip()
            label_line = _build_label_line(description, stock_code, product_description)
            return {
                'set_code': material_code,
                'stock_code': stock_code,
                'label_line': label_line,
            }
    return None


def lookup_label_data(cursor, stock_code, product_description=None):
    set_code = _find_set_code_by_stock_code(cursor, stock_code)
    if set_code:
        label_line = _find_label_line_from_dashboard(cursor, set_code, stock_code)
        if label_line:
            if product_description:
                label_line = _build_label_line(None, stock_code, product_description)
            return {
                'set_code': set_code,
                'stock_code': stock_code,
                'label_line': label_line,
            }

    search_keys = [
        stock_code,
        stock_code.replace('|', 'I'),
        stock_code.replace('|', 'l'),
        stock_code.replace('|', ''),
    ]
    seen_keys = set()
    for key in search_keys:
        if not key or key in seen_keys:
            continue
        seen_keys.add(key)
        # 정확히 _StockCode 로 끝나는 제품명을 우선 (Y2991 → P50100004)
        cursor.execute(
            '''
            SELECT product_code, product_name
            FROM export_quotation_new_products
            WHERE product_name LIKE ?
            ORDER BY
                CASE
                    WHEN UPPER(product_name) LIKE UPPER(?) THEN 0
                    ELSE 1
                END,
                LENGTH(COALESCE(product_code, '')) DESC,
                updated_at DESC
            LIMIT 1
            ''',
            (f'%{key}%', f'%_{key}'),
        )
        row = cursor.fetchone()
        if row and row[0] and row[1]:
            product_code = row[0]
            product_name = str(row[1]).strip()
            label_line = _build_label_line(product_name, stock_code, product_description)
            return {
                'set_code': product_code,
                'stock_code': stock_code,
                'label_line': label_line,
            }

    part_hit = _find_label_from_part_prices(cursor, stock_code, product_description)
    if part_hit:
        return part_hit

    # 대시보드에 없어도 템플릿 Product Description이 있으면 라벨 라인 구성
    desc = (product_description or '').strip()
    if desc and stock_code:
        return {
            'set_code': '',
            'stock_code': stock_code,
            'label_line': _build_label_line(None, stock_code, desc),
        }

    return None


def build_labels_from_template(cursor, filepath):
    parsed_entries = parse_template_entries(filepath)
    if not parsed_entries:
        return [], '업로드된 워드 파일에서 Stock Code를 찾을 수 없습니다.'

    labels = []
    missing = []
    pallet_groups_emitted = set()
    for entry in parsed_entries:
        data = lookup_label_data(
            cursor,
            entry['stock_code'],
            product_description=entry.get('product_description'),
        )
        if not data or not data.get('set_code'):
            missing.append(entry['stock_code'])
            continue
        group_id = entry.get('pallet_group_id')
        include_pallet = group_id not in pallet_groups_emitted
        if group_id is not None:
            pallet_groups_emitted.add(group_id)
        labels.append({
            **entry,
            **data,
            'include_pallet_page': include_pallet,
        })

    if not labels:
        missing_text = ', '.join(missing[:8])
        return [], f'Stock Code({missing_text})에 해당하는 세트코드를 대시보드에서 찾을 수 없습니다.'

    return labels, None


def _copy_section_layout(source_doc, target_doc):
    if not source_doc.sections or not target_doc.sections:
        return
    src = source_doc.sections[0]
    dst = target_doc.sections[0]
    dst.page_width = src.page_width
    dst.page_height = src.page_height
    dst.top_margin = src.top_margin
    dst.bottom_margin = src.bottom_margin
    dst.left_margin = src.left_margin
    dst.right_margin = src.right_margin


def _strip_tblp_pr(element):
    """페이지 고정(floating) 테이블 속성 제거 — 2페이지 빈 프레임 방지"""
    if element.tag != qn('w:tbl'):
        return
    tbl_pr = element.find(qn('w:tblPr'))
    if tbl_pr is None:
        return
    tblp_pr = tbl_pr.find(qn('w:tblpPr'))
    if tblp_pr is not None:
        tbl_pr.remove(tblp_pr)


def _trim_trailing_empty_paragraphs_in_table(element):
    """테이블 마지막 행의 빈 문단 제거"""
    if element.tag != qn('w:tbl'):
        return
    rows = [child for child in element if child.tag == qn('w:tr')]
    if not rows:
        return
    last_row = rows[-1]
    for tc in last_row.findall('./' + qn('w:tc')):
        paragraphs = [p for p in tc.findall('./' + qn('w:p'))]
        for paragraph in reversed(paragraphs):
            if _paragraph_has_visible_content(paragraph):
                break
            tc.remove(paragraph)


def _prepare_pallet_page_elements(page_elements):
    """pallet 페이지용 — 테이블 + Pallet 문단만 포함 (빈 문단·고정배치 제거)"""
    prepared = []
    for element in page_elements or []:
        if element.tag == qn('w:tbl'):
            if not any((t.text or '').strip() for t in element.findall('.//' + qn('w:t'))):
                continue
            cloned = deepcopy(element)
            _strip_tblp_pr(cloned)
            _trim_trailing_empty_paragraphs_in_table(cloned)
            prepared.append(cloned)
        elif element.tag == qn('w:p') and _paragraph_has_visible_content(element):
            text = ''.join((t.text or '') for t in element.findall('.//' + qn('w:t')))
            if PALLET_PATTERN.search(text):
                cloned = deepcopy(element)
                _remove_page_breaks_from_element(cloned)
                prepared.append(cloned)
    return prepared


DEFAULT_LABEL_TABLE_WIDTH_DXA = 14142
DEFAULT_LABEL_FRAME_HEIGHT_IN = 4.5
DEFAULT_LABEL_TOP_OFFSET_TWIPS = 1935
LABEL_BORDER_SZ = '18'


def _extract_template_table_width(template_filepath):
    if not template_filepath:
        return DEFAULT_LABEL_TABLE_WIDTH_DXA
    template_doc = Document(template_filepath)
    if not template_doc.tables:
        return DEFAULT_LABEL_TABLE_WIDTH_DXA
    tbl_pr = template_doc.tables[0]._tbl.find(qn('w:tblPr'))
    if tbl_pr is None:
        return DEFAULT_LABEL_TABLE_WIDTH_DXA
    tbl_w = tbl_pr.find(qn('w:tblW'))
    if tbl_w is None:
        return DEFAULT_LABEL_TABLE_WIDTH_DXA
    return int(tbl_w.get(qn('w:w'), str(DEFAULT_LABEL_TABLE_WIDTH_DXA)))


def _extract_template_label_top_offset(template_filepath):
    if not template_filepath:
        return DEFAULT_LABEL_TOP_OFFSET_TWIPS
    template_doc = Document(template_filepath)
    if not template_doc.tables:
        return DEFAULT_LABEL_TOP_OFFSET_TWIPS
    tbl_pr = template_doc.tables[0]._tbl.find(qn('w:tblPr'))
    if tbl_pr is None:
        return DEFAULT_LABEL_TOP_OFFSET_TWIPS
    tblp_pr = tbl_pr.find(qn('w:tblpPr'))
    if tblp_pr is None:
        return DEFAULT_LABEL_TOP_OFFSET_TWIPS
    return int(tblp_pr.get(qn('w:tblpY'), str(DEFAULT_LABEL_TOP_OFFSET_TWIPS)))


def _ensure_tbl_pr(table):
    tbl_pr = table._tbl.find(qn('w:tblPr'))
    if tbl_pr is None:
        tbl_pr = OxmlElement('w:tblPr')
        table._tbl.insert(0, tbl_pr)
    return tbl_pr


def _set_table_width_dxa(table, width_dxa):
    tbl_pr = _ensure_tbl_pr(table)
    tbl_w = tbl_pr.find(qn('w:tblW'))
    if tbl_w is None:
        tbl_w = OxmlElement('w:tblW')
        tbl_pr.append(tbl_w)
    tbl_w.set(qn('w:w'), str(width_dxa))
    tbl_w.set(qn('w:type'), 'dxa')


def _set_table_alignment(table, align='center'):
    tbl_pr = _ensure_tbl_pr(table)
    jc = tbl_pr.find(qn('w:jc'))
    if jc is None:
        jc = OxmlElement('w:jc')
        tbl_pr.append(jc)
    jc.set(qn('w:val'), align)


def _set_cell_vertical_alignment(cell, align='center'):
    tc_pr = cell._tc.get_or_add_tcPr()
    v_align = tc_pr.find(qn('w:vAlign'))
    if v_align is None:
        v_align = OxmlElement('w:vAlign')
        tc_pr.append(v_align)
    v_align.set(qn('w:val'), align)


def _set_cell_borders(cell, border_sz=LABEL_BORDER_SZ):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_borders = tc_pr.find(qn('w:tcBorders'))
    if tc_borders is None:
        tc_borders = OxmlElement('w:tcBorders')
        tc_pr.append(tc_borders)
    for edge in ('top', 'left', 'bottom', 'right'):
        element = tc_borders.find(qn(f'w:{edge}'))
        if element is None:
            element = OxmlElement(f'w:{edge}')
            tc_borders.append(element)
        element.set(qn('w:val'), 'single')
        element.set(qn('w:sz'), border_sz)
        element.set(qn('w:space'), '0')
        element.set(qn('w:color'), 'auto')


def _add_label_top_spacer(doc, top_offset_twips):
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_before = Pt(top_offset_twips / 20)
    spacer.paragraph_format.space_after = Pt(0)


def _append_page_elements(doc, page_elements):
    body = doc.element.body
    insert_at = len(body)
    if len(body) > 0 and body[-1].tag == qn('w:sectPr'):
        insert_at -= 1
    for element in page_elements:
        cloned = deepcopy(element)
        _remove_page_breaks_from_element(cloned)
        body.insert(insert_at, cloned)
        insert_at += 1


def _append_document_body(target_doc, source_doc):
    """원본 워드 본문을 그대로 복사 (업로드 템플릿 레이아웃 유지)"""
    source_children = list(source_doc.element.body)
    if source_children and source_children[-1].tag == qn('w:sectPr'):
        source_children = source_children[:-1]
    if source_children:
        _append_page_elements(target_doc, source_children)


def _append_upload_template_page(doc, entry, template_filepath):
    """Type A — 업로드 pallet 페이지 1장만 추가"""
    page_elements = _prepare_pallet_page_elements(entry.get('page_elements'))
    if page_elements:
        _append_page_elements(doc, page_elements)
        if entry.get('needs_row_update'):
            _update_template_page_row(doc, entry)
        return

    if not template_filepath:
        return

    template_doc = Document(template_filepath)
    if entry.get('needs_row_update'):
        _update_template_page_row(template_doc, entry)
    for element in _prepare_pallet_page_elements(list(template_doc.element.body)):
        _append_page_elements(doc, [element])


def _set_cell_text(cell, text):
    cell.text = _sanitize_docx_text(text)


def _update_template_page_row(doc, entry):
    for table in doc.tables:
        stock_col = qty_col = sr_col = desc_col = None
        header_row_idx = None

        for ri, row in enumerate(table.rows):
            cells = [cell.text.strip() for cell in row.cells]
            joined = ' '.join(cells).lower()
            if 'stock code' in joined or ('stock' in joined and 'code' in joined):
                header_row_idx = ri
                for ci, text in enumerate(cells):
                    lower = text.lower()
                    if lower == 'sr':
                        sr_col = ci
                    elif 'stock' in lower and 'code' in lower:
                        stock_col = ci
                    elif 'product' in lower and 'description' in lower:
                        desc_col = ci
                    elif 'quantity' in lower or lower in ('qty', "q'ty"):
                        qty_col = ci
                break

        if header_row_idx is None or stock_col is None:
            continue

        for row in table.rows[header_row_idx + 1:]:
            cells = row.cells
            if len(cells) <= stock_col:
                continue
            current = cells[stock_col].text.strip()
            if current and STOCK_CODE_PATTERN.match(current):
                if sr_col is not None and len(cells) > sr_col:
                    _set_cell_text(cells[sr_col], entry.get('sr', '1'))
                _set_cell_text(cells[stock_col], entry.get('stock_code', ''))
                if desc_col is not None and len(cells) > desc_col and entry.get('product_description'):
                    _set_cell_text(cells[desc_col], entry.get('product_description', ''))
                if qty_col is not None and len(cells) > qty_col:
                    _set_cell_text(cells[qty_col], entry.get('quantity', '1'))
                return


def _set_paragraph_spacing(paragraph, before_pt=0, after_pt=0):
    paragraph.paragraph_format.space_before = Pt(before_pt)
    paragraph.paragraph_format.space_after = Pt(after_pt)


def _set_run_font(run, size_pt, font_name=LABEL_FONT_NAME):
    run.font.name = font_name
    run.font.size = Pt(size_pt)
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.get_or_add_rFonts()
    r_fonts.set(qn('w:ascii'), font_name)
    r_fonts.set(qn('w:hAnsi'), font_name)
    r_fonts.set(qn('w:eastAsia'), font_name)
    r_fonts.set(qn('w:cs'), font_name)


def _add_center_run(paragraph, text, *, bold=False, size_pt=16, font_name=LABEL_FONT_NAME):
    run = paragraph.add_run(_sanitize_docx_text(text))
    if bold:
        run.bold = True
    _set_run_font(run, size_pt, font_name)
    return run


def _append_generated_label(doc, label, template_filepath=None):
    """생성 라벨(Type B) — 캡처와 동일한 네모 박스 1페이지"""
    _add_label_top_spacer(doc, _extract_template_label_top_offset(template_filepath))

    table = doc.add_table(rows=1, cols=1)
    _set_table_width_dxa(table, _extract_template_table_width(template_filepath))
    _set_table_alignment(table, 'center')

    row = table.rows[0]
    row.height = Inches(DEFAULT_LABEL_FRAME_HEIGHT_IN)
    cell = row.cells[0]
    _set_cell_borders(cell)
    _set_cell_vertical_alignment(cell, 'center')

    p_set = cell.paragraphs[0]
    p_set.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_paragraph_spacing(p_set, before_pt=24, after_pt=48)
    _add_center_run(p_set, label.get('set_code', ''), bold=True, size_pt=SET_CODE_FONT_SIZE_PT)

    p_line = cell.add_paragraph()
    p_line.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_paragraph_spacing(p_line, before_pt=48, after_pt=72)
    _add_center_run(p_line, label.get('label_line', ''), size_pt=DETAIL_FONT_SIZE_PT)

    p_qty = cell.add_paragraph()
    p_qty.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_paragraph_spacing(p_qty, before_pt=72, after_pt=24)
    qty = _sanitize_docx_text(label.get('quantity', '1')) or '1'
    _add_center_run(p_qty, f"Q'TY: {qty}", size_pt=DETAIL_FONT_SIZE_PT)


def generate_label_docx(labels, template_filepath=None):
    """
    pallet 페이지 + Stock Code별 P코드 라벨 페이지 생성
    - pallet 테이블에 여러 Stock Code가 있으면: pallet 1장 + 각 코드별 라벨 페이지
    - 예: Y0149, H3571|00001 → 1p pallet + 2p Y0149라벨 + 3p H3571라벨
    """
    if not labels:
        doc = Document()
        buffer = io.BytesIO()
        doc.save(buffer)
        buffer.seek(0)
        return buffer.getvalue()

    doc = Document()
    if template_filepath:
        _copy_section_layout(Document(template_filepath), doc)

    first_content = True
    for entry in labels:
        if entry.get('include_pallet_page', True):
            if not first_content:
                doc.add_page_break()
            _append_upload_template_page(doc, entry, template_filepath)
            first_content = False

        if not first_content:
            doc.add_page_break()
        _append_generated_label(doc, entry, template_filepath)
        first_content = False

    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()
